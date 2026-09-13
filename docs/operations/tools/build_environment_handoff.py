"""Create the Word handoff from its reviewed Markdown source with bundled docx.

Run using Codex's bundled document Python, not the business Conda environment.
Only reads the public report; never reads local credentials or scan bodies.
"""

from pathlib import Path
import re

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[3]
SOURCE = ROOT / "docs/operations/v1-environment-acceptance-20260913.md"
OUTPUT = ROOT / "_local_artifacts/environment/deliverables/V1开发环境验收与行为说明_20260913.docx"


def clean(text):
    return re.sub(r"`([^`]+)`", r"\1", text).replace("**", "")


def main():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.7)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "List Bullet"):
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(11)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.15
    doc.styles["Title"].font.size = Pt(22)
    doc.styles["Heading 1"].font.size = Pt(15)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(13)
    doc.styles["Heading 2"].font.size = Pt(12)
    doc.core_properties.title = "V1 开发环境验收与行为说明"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    index = 0
    code = False
    while index < len(lines):
        line = lines[index]
        index += 1
        if not line.strip():
            continue
        if line.startswith("```"):
            code = not code
            continue
        if line.startswith("|"):
            rows = [line]
            while index < len(lines) and lines[index].startswith("|"):
                rows.append(lines[index])
                index += 1
            values = [[clean(c.strip()) for c in row.strip("|").split("|")]
                      for row in rows if not re.match(r"^\|[\s|:-]+\|$", row)]
            table = doc.add_table(rows=0, cols=len(values[0]))
            table.autofit = False
            widths = [1.55, 2.55, 3.0] if "层次" in values[0] else [1.95, 3.55, 1.6]
            if "检查" in values[0]:
                widths = [1.9, 1.3, 3.9]
            for col, width in zip(table.columns, widths):
                col.width = Inches(width)
            borders = OxmlElement("w:tblBorders")
            for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
                edge = OxmlElement("w:" + side)
                for attr, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                    edge.set(qn("w:" + attr), value)
                borders.append(edge)
            table._tbl.tblPr.append(borders)
            for n, row in enumerate(values):
                cells = table.add_row().cells
                properties = cells[0]._tc.getparent().get_or_add_trPr()
                properties.append(OxmlElement("w:cantSplit"))
                if n == 0:
                    properties.append(OxmlElement("w:tblHeader"))
                for cell, text, width in zip(cells, row, widths):
                    cell.width = Inches(width)
                    cell.vertical_alignment = 1
                    tc = cell._tc.get_or_add_tcPr()
                    margins = OxmlElement("w:tcMar")
                    for side in ("top", "left", "bottom", "right"):
                        edge = OxmlElement("w:" + side)
                        edge.set(qn("w:w"), "90")
                        edge.set(qn("w:type"), "dxa")
                        margins.append(edge)
                    tc.append(margins)
                    paragraph = cell.paragraphs[0]
                    if values[0][-1] == "GitHub" and cell is cells[-1]:
                        paragraph.alignment = 1
                    paragraph.paragraph_format.space_after = Pt(2)
                    paragraph.paragraph_format.space_before = Pt(2)
                    paragraph.paragraph_format.line_spacing = 1.05
                    run = paragraph.add_run(text)
                    run.font.size = Pt(10.5)
                    if n == 0:
                        run.bold = True
                        shade = OxmlElement("w:shd")
                        shade.set(qn("w:fill"), "E8EEF4")
                        tc.append(shade)
            doc.add_paragraph().paragraph_format.space_after = Pt(1)
        elif line.startswith("# "):
            doc.add_paragraph(clean(line[2:]), style="Title")
        elif line.startswith("## "):
            doc.add_paragraph(clean(line[3:]), style="Heading 1")
        elif line.startswith("- "):
            doc.add_paragraph(clean(line[2:]), style="List Bullet")
        else:
            paragraph = doc.add_paragraph(clean(line))
            if code:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.runs[0].font.size = Pt(10)
    # Default document themes can carry a blue Title paragraph border.
    # Remove paragraph borders only; intentional table borders are preserved.
    for tree in (doc.styles.element, doc.element):
        for border in list(tree.xpath(".//w:pBdr")):
            border.getparent().remove(border)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
